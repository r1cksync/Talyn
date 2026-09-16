"""Generate only synthetic documents and fixture media, with no customer data."""
import math
import shutil
import struct
import subprocess
import wave
from pathlib import Path

from docx import Document

root = Path(__file__).resolve().parents[1] / "frontend" / "test-fixtures"
root.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_heading("Alex Morgan — Synthetic resume", 0)
doc.add_paragraph("Synthetic test persona. No real person is represented.")
doc.add_heading("Experience", 1)
doc.add_paragraph("Built a Python API with PostgreSQL transactions and idempotency keys. Compared designs with a team and measured a 40% reduction in duplicate processing.")
doc.add_heading("Projects", 1)
doc.add_paragraph("Designed reliable background jobs, tested retries, and presented tradeoffs to stakeholders.")
doc.add_paragraph("Skills: Python, PostgreSQL, testing, collaboration. Education: fictional computer science degree.")
doc.save(root / "resume.docx")
(root / "candidates.csv").write_text("name,email\nAlex Morgan,alex@example.com\n", encoding="utf-8")
with wave.open(str(root / "audio.wav"), "wb") as output:
    output.setparams((1, 2, 48000, 0, "NONE", "not compressed"))
    output.writeframes(b"".join(struct.pack("<h", int(2000*math.sin(2*math.pi*220*i/48000))) for i in range(48000*5)))
ffmpeg = shutil.which("ffmpeg")
if not ffmpeg:
    import imageio_ffmpeg
    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
subprocess.run([ffmpeg, "-y", "-v", "error", "-f", "lavfi", "-i", "testsrc2=size=320x240:rate=15", "-t", "3", "-pix_fmt", "yuv420p", str(root / "camera.y4m")], check=True)
subprocess.run([ffmpeg, "-y", "-v", "error", "-f", "lavfi", "-i", "testsrc2=size=320x240:rate=15", "-t", "2", "-c:v", "libvpx", "-b:v", "150k", str(root / "clip.webm")], check=True)
print("Generated synthetic DOCX, CSV, WAV, Y4M, and independently playable WebM fixtures.")
